"""Vorläufige Kurz-BWA aus einer Datensicherung des Angebots-Tools.
Aufruf:  python3 kurz_bwa_aus_backup.py <Sicherung.json> [daten.json]
Schreibt Kurz-BWA_<Stand>.xlsx: Monatsübersicht Jan–Sep, Forderungen, Auftragsbestand/Forecast, Hinweise.
Keine Buchhaltungs-BWA: Grundlage sind Ausgangsrechnungen, zugeordnete Eingangsrechnungen und der Overhead-Plan der App."""
import json, sys, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
DATEN = sys.argv[2] if len(sys.argv) > 2 else None
raw = json.load(open(SRC)); db = raw["data"]
stand_iso = (raw.get("_date") or datetime.date.today().isoformat())[:10]
stand = ".".join(reversed(stand_iso.split("-")))
JAHR = stand_iso[:4]
MONATE = [f"{JAHR}-{m:02d}" for m in range(1, int(stand_iso[5:7]) + 1)]
MLBL = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
SPARTEN = {"marketing": "Marketing", "auto": "Fahrzeugkosten (Auto-Abo)", "auto_strom": "Fahrzeugkosten Strom", "reise": "Reisekosten",
           "betriebskosten": "Betriebskosten (Hosting, Software, Büro)", "geschaeftsausstattung": "Geschäftsausstattung", "versicherungen": "Versicherungen",
           "mobilfunk": "Mobilfunk", "messe": "Messen", "allgemein": "Allgemein / Sonstiges"}
def mon(s):
    s = str(s or ""); return s[:7] if len(s) >= 7 and s[4] == "-" else ""
def num(x):
    try: return float(x or 0)
    except: return 0.0

P = db["projects"]
R = db["rechnungen"] + db.get("rechManual", [])
# Umsatz: Ausgangsrechnungen netto nach Rechnungsdatum
ums = collections.Counter(); offen = []
for p in P:
    for a in p.get("ausgang") or []:
        ums[mon(a.get("datum"))] += num(a.get("amount"))
        if not a.get("paid"): offen.append((p.get("nummer"), p.get("name"), a.get("label"), a.get("datum"), num(a.get("amount"))))
# Projektkosten (Eingangsrechnungen, Projekt zugeordnet) und Betriebsausgaben (allgemein, je Sparte), netto nach Rechnungsdatum
proj = collections.Counter(); gen = collections.defaultdict(collections.Counter); neu = []
for r in R:
    mk = mon(r.get("invoiceDate") or r.get("received")); amt = num(r.get("amount"))
    if r.get("status") == "zugewiesen" and r.get("assignType") == "project": proj[mk] += amt
    elif r.get("status") == "zugewiesen" and r.get("assignType") == "general": gen[r.get("generalCat") or "allgemein"][mk] += amt
    elif r.get("status") == "neu": neu.append((mk, r.get("senderName") or "", r.get("subject") or "", amt))
# Personal: Plan aus dem Overhead (Gehalt Auszahlung + Lohnsteuer/Soli), monatlich ab „von“
fix = db.get("overheadFix") or []
def fix_month(mk, sparten):
    s = 0.0
    for f in fix:
        if f.get("sparte") not in sparten or f.get("art") == "einmalig" or not f.get("von"): continue
        if mk < f["von"] or (f.get("bis") and mk > f["bis"]): continue
        s += num(f.get("betrag"))
    return s
personal = {mk: fix_month(mk, {"gehalt", "finanzamt"}) for mk in MONATE}
planfix = {mk: fix_month(mk, set(SPARTEN)) for mk in MONATE}
# Forecast / Auftragsbestand
fc = None
if DATEN:
    try: fc = json.load(open(DATEN)).get("forecast")
    except Exception: fc = None
best = sorted([(p.get("nummer"), p.get("name"), p.get("ampel")) for p in P if p.get("nummer") and not p.get("parentNr")], key=lambda x: str(x[0]))

# ---------- Workbook ----------
wb = Workbook()
F = "Arial"
H1 = Font(name=F, size=14, bold=True); H2 = Font(name=F, size=10, bold=True); N = Font(name=F, size=10); B = Font(name=F, size=10, bold=True)
BLUE = Font(name=F, size=10, color="0000FF"); GREY = Font(name=F, size=9, italic=True, color="666666")
FILL = PatternFill("solid", fgColor="F4EEE2"); FILL2 = PatternFill("solid", fgColor="EFECE7")
thin = Side(style="thin", color="C8C8C8"); TOP = Border(top=Side(style="thin", color="3D3D3D")); DBL = Border(top=Side(style="double", color="3D3D3D"))
EUR = '#,##0 €;[Red]-#,##0 €;-'; PCT = '0.0%;[Red]-0.0%;-'

ws = wb.active; ws.title = "Kurz-BWA 2026"
ws["A1"] = f"NOVA WORKS GmbH · Vorläufige Kurz-BWA {MLBL[0]}–{MLBL[len(MONATE)-1]} {JAHR}"; ws["A1"].font = H1
ws["A2"] = f"Stand der Datensicherung: {stand} · alle Beträge netto in € · Grundlage: Nova Works App (Ausgangsrechnungen, zugeordnete Eingangsrechnungen, Overhead-Plan) · keine Buchhaltungs-BWA"; ws["A2"].font = GREY
hdr_row = 4
ws.cell(hdr_row, 1, "Position").font = H2
for i, mk in enumerate(MONATE): c = ws.cell(hdr_row, 2 + i, f"{MLBL[int(mk[5:7]) - 1]} {JAHR[2:]}"); c.font = H2; c.alignment = Alignment(horizontal="right")
cS = 2 + len(MONATE); ws.cell(hdr_row, cS, f"Summe {JAHR}").font = H2; ws.cell(hdr_row, cS).alignment = Alignment(horizontal="right")
ws.cell(hdr_row, cS + 1, "Anteil am Umsatz").font = H2
for c in range(1, cS + 2): ws.cell(hdr_row, c).fill = FILL; ws.cell(hdr_row, c).border = Border(bottom=thin)
L = {}
row = hdr_row + 1
def put(label, values, font=N, fmt=EUR, bold=False, border=None, key=None, comment=None, pct_of=None):
    global row
    ws.cell(row, 1, label).font = B if bold else font
    if values is not None:
        for i, mk in enumerate(MONATE):
            c = ws.cell(row, 2 + i, round(values.get(mk, 0.0), 2) if isinstance(values, dict) else values); c.number_format = fmt; c.font = BLUE if font is BLUE else (B if bold else N)
    rng = f"{get_column_letter(2)}{row}:{get_column_letter(1 + len(MONATE))}{row}"
    s = ws.cell(row, cS, f"=SUM({rng})"); s.number_format = fmt; s.font = B if bold else N
    if pct_of:
        pc = ws.cell(row, cS + 1, f'=IF({get_column_letter(cS)}{pct_of}=0,0,{get_column_letter(cS)}{row}/{get_column_letter(cS)}{pct_of})'); pc.number_format = PCT; pc.font = GREY
    if border:
        for c in range(1, cS + 2): ws.cell(row, c).border = border
    if comment: ws.cell(row, 1).comment = Comment(comment, "Nova Works App")
    if key: L[key] = row
    row += 1
def formula_row(label, fn, key=None, bold=False, border=None, fmt=EUR, pct_of=None, fill=None):
    global row
    ws.cell(row, 1, label).font = B if bold else N
    for i in range(len(MONATE) + 1):
        col = get_column_letter(2 + i); c = ws.cell(row, 2 + i, fn(col)); c.number_format = fmt; c.font = B if bold else N
    if pct_of:
        pc = ws.cell(row, cS + 1, f'=IF({get_column_letter(cS)}{pct_of}=0,0,{get_column_letter(cS)}{row}/{get_column_letter(cS)}{pct_of})'); pc.number_format = PCT; pc.font = GREY
    if border or fill:
        for c in range(1, cS + 2):
            if border: ws.cell(row, c).border = border
            if fill: ws.cell(row, c).fill = fill
    if key: L[key] = row
    row += 1

put("Umsatzerlöse (Ausgangsrechnungen, netto)", ums, font=BLUE, bold=True, key="ums", comment="Summe der Ausgangsrechnungen der Projekte nach Rechnungsdatum (Nova Works App → Projekt → Ausgangsrechnungen, aus Lexware Office synchronisiert).")
put("Projektkosten / Fremdleistungen (Eingangsrechnungen, Projekt zugeordnet)", proj, font=BLUE, key="proj", comment="Eingangsrechnungen mit Status „zugewiesen“ und Zuordnung zu einem Projekt, netto nach Rechnungsdatum (Nova Works App → Rechnungen). Enthält Fremdpersonal, Anmietung, Transport und Sonstiges; eine Aufteilung nach Kostenart ist in der App nicht hinterlegt.")
formula_row("Rohertrag", lambda c: f"={c}{L['ums']}-{c}{L['proj']}", key="roh", bold=True, border=TOP, pct_of=L["ums"])
row += 1
put("Personalkosten GF (Gehalt Auszahlung + Lohnsteuer/Soli, Plan)", personal, font=BLUE, key="pers", comment="Planwerte aus Nova Works App → Overhead → Fixkosten: Gehalt Kilian Körber (Auszahlung) 4.958,38 €/Monat + Lohnsteuer/Soli 2.041,62 €/Monat ab Februar 2026. Arbeitgeberanteile zur Sozialversicherung sind nicht erfasst – die Buchhaltungs-BWA weist höhere Personalkosten aus.")
ws.cell(row, 1, "Betriebsausgaben (Eingangsrechnungen allgemein, je Sparte)").font = H2; row += 1
first_sp = row
for slug, lbl in SPARTEN.items():
    put("   " + lbl, gen.get(slug, {}), font=BLUE, key="sp_" + slug)
last_sp = row - 1
formula_row("Summe Betriebsausgaben", lambda c: f"=SUM({c}{first_sp}:{c}{last_sp})", key="bak", bold=True, border=TOP)
row += 1
formula_row("Betriebsergebnis vor Abschreibungen, Zinsen und Steuern", lambda c: f"={c}{L['roh']}-{c}{L['pers']}-{c}{L['bak']}", key="erg", bold=True, border=DBL, pct_of=L["ums"], fill=FILL2)
row += 1
ws.cell(row, 1, "Nachrichtlich").font = H2; row += 1
put("Fixkosten laut Overhead-Plan (ohne Personal, Vergleichswert)", planfix, font=GREY, fmt=EUR, key="planfix", comment="Monatliche Fixkosten laut Nova Works App → Overhead (ohne Gehalt/Finanzamt, ohne einmalige Messekosten). Dient nur dem Abgleich mit den Ist-Betriebsausgaben.")
neu_m = collections.Counter(); [neu_m.__setitem__(m_, neu_m[m_] + a) for m_, _, _, a in neu]
put("Noch nicht zugeordnete Eingangsrechnungen (Status „neu“)", neu_m, font=GREY, key="neu", comment="Eingangsrechnungen im Postfach der App ohne Zuordnung. Sie fehlen oben in den Projektkosten bzw. Betriebsausgaben.")
formula_row("Kumulierter Umsatz", lambda c: (f"=SUM($B${L['ums']}:{c}{L['ums']})" if c != get_column_letter(cS) else f"={c}{L['ums']}"), key="kum", fmt=EUR)
row += 1
ws.cell(row, 1, "Legende: blaue Zahlen = Werte aus der Datensicherung · schwarze Zahlen = Formeln · graue Zeilen = nachrichtlich. Monatszuordnung nach Rechnungsdatum (nicht nach Zahlung oder Leistungszeitraum).").font = GREY
ws.column_dimensions["A"].width = 62
for i in range(len(MONATE) + 1): ws.column_dimensions[get_column_letter(2 + i)].width = 13
ws.column_dimensions[get_column_letter(cS + 1)].width = 16
ws.freeze_panes = ws.cell(hdr_row + 1, 2)

# ---------- Forderungen ----------
w2 = wb.create_sheet("Offene Forderungen")
w2["A1"] = f"Offene Ausgangsrechnungen (noch nicht bezahlt) · Stand {stand} · netto"; w2["A1"].font = H1
for j, h in enumerate(["Projekt", "Projektname", "Rechnung", "Rechnungsdatum", "Betrag netto €", "Tage offen"]):
    c = w2.cell(3, 1 + j, h); c.font = H2; c.fill = FILL
heute = datetime.date.fromisoformat(stand_iso)
r0 = 4
for i, (nr, name, lbl, dat, amt) in enumerate(sorted(offen, key=lambda x: (x[3] or "", x[0] or ""))):
    rr = r0 + i
    w2.cell(rr, 1, nr).font = N; w2.cell(rr, 2, name).font = N; w2.cell(rr, 3, lbl).font = N
    w2.cell(rr, 4, dat).font = N
    c = w2.cell(rr, 5, round(amt, 2)); c.number_format = EUR; c.font = BLUE
    try: w2.cell(rr, 6, (heute - datetime.date.fromisoformat(dat)).days).font = N
    except Exception: pass
rend = r0 + len(offen) - 1
w2.cell(rend + 1, 4, "Summe").font = B
c = w2.cell(rend + 1, 5, f"=SUM(E{r0}:E{rend})"); c.number_format = EUR; c.font = B; c.border = TOP
for col, wdt in zip("ABCDEF", (10, 34, 24, 16, 16, 11)): w2.column_dimensions[col].width = wdt
w2.freeze_panes = "A4"

# ---------- Auftragsbestand / Forecast ----------
w3 = wb.create_sheet("Auftragsbestand")
w3["A1"] = f"Auftragsbestand und Forecast · Nova Works App · Stand {stand} · netto"; w3["A1"].font = H1
rr = 3
if fc:
    for key, lbl in (("zwoelf", "Nächste 12 Monate"), ("jahr", "Kalenderjahr " + JAHR), ("naechstes", "Kalenderjahr " + str(int(JAHR) + 1))):
        X = fc.get(key)
        if not X: continue
        T = X["T"]
        w3.cell(rr, 1, f"{lbl} ({X.get('periodLbl','')})").font = H2; rr += 1
        for l2, v, cm in (("Bestätigt / fest (Projektsummen)", T["conf"], "Projektsumme bestätigter Projekte und feste Planeinträge"),
                          ("Offen (Angebote auf Anfrage, ungewichtet)", T["open"], "Angebots-Netto der Projekte „auf Anfrage“"),
                          ("Offen gewichtet (× Wahrscheinlichkeit)", T["wo"], "Standard-Wahrscheinlichkeit laut App"),
                          ("Forecast Umsatz (bestätigt + offen gewichtet)", T["wc"] + T["wo"], ""),
                          ("Bereits abgerechnet", T["erl"], "Ausgangsrechnungen im Zeitraum"),
                          ("Geplante Projektkosten", T["soll"], "Kalkulations-Soll bzw. Angebots-EK"),
                          ("Marge (Forecast)", T["marge"], ""), ("Fixkosten (Overhead-Plan)", T["fix"], ""), ("Ergebnis nach Fixkosten", T["erg"], "")):
            w3.cell(rr, 1, "   " + l2).font = N; c = w3.cell(rr, 2, round(v, 2)); c.number_format = EUR; c.font = BLUE
            if cm: w3.cell(rr, 3, cm).font = GREY
            rr += 1
        rr += 1
    Z = fc["zwoelf"]
    w3.cell(rr, 1, "Umsatz je Monat, nächste 12 Monate").font = H2; rr += 1
    for j, h in enumerate(["Monat", "Bestätigt / fest", "Offen gewichtet", "Abgerechnet", "Forecast"]):
        c = w3.cell(rr, 1 + j, h); c.font = H2; c.fill = FILL
    rr += 1; m0 = rr
    for m in Z["monate"]:
        w3.cell(rr, 1, m["label"]).font = N
        for j, k in enumerate(("bestaetigt", "offen", "abgerechnet")):
            c = w3.cell(rr, 2 + j, round(m[k], 2)); c.number_format = EUR; c.font = BLUE
        c = w3.cell(rr, 5, f"=B{rr}+C{rr}"); c.number_format = EUR; c.font = N
        rr += 1
    w3.cell(rr, 1, "Summe").font = B
    for j in range(4):
        col = get_column_letter(2 + j); c = w3.cell(rr, 2 + j, f"=SUM({col}{m0}:{col}{rr-1})"); c.number_format = EUR; c.font = B; c.border = TOP
    rr += 2
else:
    w3.cell(rr, 1, "Kein Forecast übergeben (daten.json fehlt).").font = GREY; rr += 2
w3.cell(rr, 1, f"Projekte {JAHR} nach Status").font = H2; rr += 1
cnt = collections.Counter(a for _, _, a in best)
for a, lbl in (("bestaetigt", "bestätigt"), ("anfrage", "angeboten / auf Anfrage"), ("abgesagt", "abgesagt")):
    w3.cell(rr, 1, "   " + lbl).font = N; w3.cell(rr, 2, cnt.get(a, 0)).font = BLUE; rr += 1
w3.column_dimensions["A"].width = 48; w3.column_dimensions["B"].width = 18; w3.column_dimensions["C"].width = 18; w3.column_dimensions["D"].width = 16; w3.column_dimensions["E"].width = 16

# ---------- Hinweise ----------
w4 = wb.create_sheet("Hinweise")
w4["A1"] = "Hinweise zu dieser Auswertung"; w4["A1"].font = H1
hints = [
    ("Was das ist", f"Eine vorläufige Kurz-BWA aus den Betriebsdaten der Nova Works App (Datensicherung vom {stand}). Sie zeigt Umsatz, Projektkosten, Personal und Betriebsausgaben je Monat nach Rechnungsdatum."),
    ("Was das nicht ist", "Keine Betriebswirtschaftliche Auswertung aus der Finanzbuchhaltung. Es fehlen Abschreibungen, Zinsen, Steuern, Arbeitgeberanteile zur Sozialversicherung, Abgrenzungen und Bestandsveränderungen. Für Creditreform ist die BWA aus Lexware Office bzw. vom Steuerberater maßgeblich."),
    ("Umsatz", "Ausgangsrechnungen der Projekte, netto, nach Rechnungsdatum. Zusätze ohne Rechnung und Projekte ohne Rechnungsstellung sind nicht enthalten."),
    ("Projektkosten", "Eingangsrechnungen mit Projektzuordnung, netto. Die App weist in der Auswertung „Kosten (Ist)“ nur Rechnungen mit Positionszuordnung aus; hier sind alle projektbezogenen Rechnungen enthalten, deshalb kann der Wert leicht über „Kosten (Ist)“ liegen."),
    ("Personal", "Nur Geschäftsführergehalt (Auszahlung) plus Lohnsteuer/Soli laut Overhead-Plan. Keine Sozialversicherung, keine weiteren Mitarbeiter."),
    ("Betriebsausgaben", "Eingangsrechnungen ohne Projektbezug nach Sparte der App. Monatliche Fixkosten, zu denen noch keine Rechnung erfasst ist, fehlen; der Overhead-Plan steht zum Vergleich nachrichtlich darunter."),
    ("Nicht zugeordnet", "Eingangsrechnungen mit Status „neu“ sind weder in den Projektkosten noch in den Betriebsausgaben enthalten (Summe siehe nachrichtliche Zeile)."),
    ("Monatszuordnung", "Nach Rechnungsdatum, nicht nach Leistungszeitraum oder Zahlungseingang. Vorkasse-Rechnungen können deshalb vor dem Veranstaltungsmonat liegen."),
    ("Was Creditreform üblicherweise erwartet", "Aktuelle BWA mit Summen- und Saldenliste (per 30.09.), Eröffnungsbilanz bzw. letzter Jahresabschluss, Gesellschafterstruktur und Stammkapital, Auftragsbestand und Umsatzplanung (siehe Blatt Auftragsbestand), Bankverbindlichkeiten und Leasingverpflichtungen, Forderungsbestand (siehe Blatt Offene Forderungen)."),
]
for i, (k, v) in enumerate(hints):
    w4.cell(3 + i, 1, k).font = B; c = w4.cell(3 + i, 2, v); c.font = N; c.alignment = Alignment(wrap_text=True, vertical="top")
    w4.row_dimensions[3 + i].height = 48
w4.column_dimensions["A"].width = 30; w4.column_dimensions["B"].width = 120

out = f"Kurz-BWA_{stand_iso}.xlsx"
wb.save(out); print("geschrieben", out)
print("Umsatz je Monat:", {k: round(v) for k, v in sorted(ums.items())}, "Summe", round(sum(ums.values())))
print("Projektkosten:", round(sum(proj.values())), "Personal:", round(sum(personal.values())), "Betriebsausgaben:", round(sum(sum(c.values()) for c in gen.values())), "neu:", round(sum(a for *_, a in neu)))
print("offene Forderungen:", len(offen), round(sum(a for *_, a in offen)))
